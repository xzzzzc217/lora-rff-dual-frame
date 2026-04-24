"""
生成"项目工作原始记录 v2"Word文档 — 结题报告附录
包含：文献调研、数据采集、特征工程、KNN初期方案、CNN频谱图实验、
DR-RFF/DeepCRF复现、集成学习框架、消融实验、问题排查等完整记录
"""

import os
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "项目工作原始记录_v2.docx")
RESULTS_DIR = r"C:\Users\21398\Desktop\sophomore\SRTP\code\ensemble_56_devices\results"

doc = Document()

# ============ 页面设置 ============
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)


# ============ 辅助函数 ============

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    return h


def add_para(text, bold=False, font_size=Pt(12), font_name="宋体",
             alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=Pt(6)):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = Pt(22)
    run = p.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.font.bold = bold
    return p


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
    run.font.size = Pt(10.5)
    run.font.bold = True
    return p


def set_cell_shading(cell, color):
    tc = cell._element.get_or_add_tcPr()
    shd = tc.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): color,
    })
    tc.append(shd)


def set_cell(cell, text, bold=False, font_size=Pt(10), color=None,
             alignment=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
    run.font.size = font_size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def fmt(v):
    return f"{float(v)*100:.1f}%"


def add_image_safe(path, width=Cm(15)):
    if os.path.exists(path):
        doc.add_picture(path, width=width)
        last_para = doc.paragraphs[-1]
        last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_para(f"[图片缺失: {os.path.basename(path)}]",
                 alignment=WD_ALIGN_PARAGRAPH.CENTER)


def make_table(headers, data, header_color="2E5A3E"):
    """通用表格构建"""
    n_rows = len(data) + 1
    n_cols = len(headers)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, header_color)
        set_cell(cell, h, bold=True, font_size=Pt(10),
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, row_data in enumerate(data):
        for j, val in enumerate(row_data):
            set_cell(table.rows[i + 1].cells[j], val, font_size=Pt(10))
    return table


def build_ablation_table(data_rows, headers, col_keys, col_widths):
    n_rows = len(data_rows) + 1
    n_cols = len(headers)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = w
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, "2E5A3E")
        set_cell(cell, h, bold=True, font_size=Pt(9),
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, rd in enumerate(data_rows):
        max_val = float(rd["最大正确率"])
        for j, key in enumerate(col_keys):
            cell = table.rows[i + 1].cells[j]
            val = rd[key]
            if key == "特征组合":
                set_cell(cell, val, font_size=Pt(9),
                         alignment=WD_ALIGN_PARAGRAPH.LEFT)
            else:
                text = fmt(val)
                is_max = abs(float(val) - max_val) < 0.0001
                is_ensemble = (key == "集成学习")
                if is_ensemble:
                    set_cell_shading(cell, "E8F5E9")
                set_cell(cell, text, bold=is_max, font_size=Pt(9))
    return table


# ===================================================================
# 正文开始
# ===================================================================

# 大标题
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_para.paragraph_format.space_after = Pt(12)
run = title_para.add_run("项目工作原始记录")
run.font.name = "黑体"
run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
run.font.size = Pt(22)
run.font.bold = True

add_para("项目名称：面向LoRa设备的连续双帧射频指纹提取与识别技术", bold=True, font_size=Pt(14))
add_para("项目编号：202510286066", font_size=Pt(12))
add_para("项目级别：国家级", font_size=Pt(12))
add_para("")

# ================================================================
# 一、文献调研与技术背景
# ================================================================
add_heading("一、文献调研与技术背景", level=1)

add_para("项目组围绕射频指纹识别（Radio Frequency Fingerprinting, RFF）开展了系统性文献调研，"
         "共阅读分析二十余篇中外高质量文献，从射频指纹识别的总体框架出发，系统梳理了"
         "物理层信号处理、深度学习建模、特征工程等方面的研究现状与技术路线。")

add_heading("1.1 射频指纹识别研究现状", level=2)

add_para("(1) 基于深度学习的方案：Robyns等人提出基于整个LoRa帧信号的神经网络分类器，"
         "针对相同芯片设备的识别准确率达59%。Shen等人研究CFO的时漂特性，"
         "通过补偿后结合频谱图输入CNN，在20个LoRa设备中实现97.61%的准确率。"
         "Al-Shawabka等人发现跨天测试时10个LoRa设备识别率仅19%，通过数据扩增提升至36%。"
         "Shen等人进一步设计仅需训练一次的射频指纹提取器，在60个设备中实现88.67%的识别率。")

add_para("(2) 基于特征工程的方案：Peng等人提出差分星座图方法消除信道影响，"
         "保留CFO和IQ不平衡特征，在16个ZigBee设备上实现90%以上准确率。"
         "Wu等人将该方法应用于LoRa，4个设备准确率达98.66%。"
         "Brik等人设计PARADIS系统，综合CFO和IQ不平衡特征，结合KNN和SVM，"
         "识别100多个802.11b网卡准确率超99%。"
         "Song等人结合KNN算法，在LoRa设备上进一步验证了特征工程方法的有效性。")

add_para("(3) 集成学习在信号分类中的优势：集成学习能够显著提升分类准确性和稳定性。"
         "单个分类器在信号分类中存在局限性——SVM对参数选择敏感、决策树易过拟合、KNN受噪声干扰大，"
         "而集成学习通过多分类器协同输出获得更可靠的置信度估计。"
         "与深度学习相比，集成学习模型复杂度较低，不易过拟合训练数据中的信道特征，"
         "天然支持开集识别与拒识机制。")

add_heading("1.2 调研结论与研究思路", level=2)

add_para("调研发现当前方案的主要局限性包括：(1) 深度学习模型泛化能力不足，跨天性能大幅下降；"
         "(2) 系统可扩展性差，设备变化需重新训练整个模型；"
         "(3) 基于特征工程的方案协议依赖性强，CFO受温度和老化影响存在长期稳定性问题。"
         "基于上述分析，本项目确定了'基于连续双帧CFO特征 + 集成学习'的技术路线，"
         "兼顾特征的可解释性和模型的鲁棒性。")

# 文献调研对比表
add_heading("1.3 核心文献对比", level=2)

make_table(
    ["方法", "类型", "设备数", "核心技术", "主要结果"],
    [
        ["Shen等[2021]", "深度学习", "20 LoRa", "CFO+频谱图+CNN", "97.61%"],
        ["Al-Shawabka等[2021]", "深度学习", "10 LoRa", "数据扩增+CNN", "跨天36%"],
        ["Shen等[2022]", "深度学习", "60 LoRa", "信道无关频谱图", "88.67%"],
        ["DR-RFF[Xie等]", "深度学习", "ZigBee", "解耦表示学习", "跨信道鲁棒"],
        ["DeepCRF[Kong等]", "深度学习", "WiFi", "CSI+对比学习", "信道鲁棒"],
        ["Brik等[PARADIS]", "特征工程", "100+ 802.11b", "CFO+IQ+KNN/SVM", ">99%"],
        ["Wu等", "特征工程", "4 LoRa", "差分星座图", "98.66%"],
        ["本项目", "特征工程+集成", "56 LoRa", "双帧CFO+相位+集成", "跨天96.13%"],
    ]
)
add_caption("表1  核心文献与本项目方法对比")

# ================================================================
# 二、数据采集与数据集建设
# ================================================================
add_heading("二、数据采集与数据集建设", level=1)

add_heading("2.1 数据采集方案", level=2)
add_para("接收端：使用USRP N210通过GNU Radio进行原始I/Q射频信号采集，"
         "采样率设置为1MSps（数据集A/B）和5MSps（数据集C），中心频率433MHz，带宽125kHz。")
add_para("发射端：基于Arduino平台完成5种型号LoRa设备的连续双帧信号发射代码编写与调试，"
         "实现各型号设备的统一配置与控制。")
add_para("采集环境：在实验室内进行视距（LoS）和非视距（NLoS）场景下的数据采集，"
         "分别在不同天进行多轮采集以构建跨天测试集。")

add_heading("2.2 数据集概况", level=2)

make_table(
    ["数据集", "设备型号", "芯片", "设备数", "采样率", "样本数(训/同日/跨日)"],
    [
        ["A", "Heltec LoRa32", "SX1262", "8", "1 MHz", "4150/2095/2227"],
        ["B", "RA01-SC", "LLCC68", "18", "1 MHz", "2745/910/930"],
        ["C", "Heltec CubeCell", "SX1262", "30", "5 MHz", "6900/3450/3450"],
    ]
)
add_caption("表2  实验数据集概况")

add_para("")
add_para("LoRa信号参数：带宽B=125kHz，扩频因子SF=7，符号周期T=128/B≈1.024ms，"
         "每帧前导码含8个upchirp，采样点数L=T×fs（1MHz下L=1024）。")

# ================================================================
# 三、信号预处理与特征工程
# ================================================================
add_heading("三、信号预处理与特征工程", level=1)

add_heading("3.1 信号预处理流程", level=2)
add_para("接收到的信号受到硬件损伤的影响，为提取设备指纹需经过以下预处理步骤：")
add_para("(1) 帧同步（Schmidl-Cox算法）：利用前导码重复upchirp的自相关特性进行粗同步，"
         "归一化自相关系数M(d)阈值设为0.95。精同步参考Robyns等人的方法，"
         "通过计算瞬时频率与理想upchirp频率的互相关确定精确起始位置。"
         "实现上采用numpy向量化（cumsum + stride_tricks），处理速度约2秒/设备。")
add_para("(2) 前导码提取：同步完成后提取开头8个upchirp符号作为设备指纹的原始信号。")
add_para("(3) 载波频偏（CFO）补偿：粗估计利用单帧瞬时频率均值，精估计利用帧间相位差。"
         "补偿后的信号用于相位特征提取。")
add_para("(4) 幅度归一化：通过信号振幅的均方根进行归一化，消除发射功率和通信距离的影响。")

add_heading("3.2 特征提取与选择", level=2)

make_table(
    ["特征类别", "维度", "计算方法", "跨天稳定性"],
    [
        ["双帧差分相位", "8", "两帧归一化相位差", "低"],
        ["CFO原始值", "2", "帧1/帧2载波频偏", "高"],
        ["CFO派生", "2", "均值、差值", "高"],
        ["RSSI", "4", "帧1/帧2信号强度及均值/差值", "低"],
        ["IQ统计", "8", "I/Q均值、标准差 × 2帧", "低"],
        ["STFT谱特征", "6", "谱质心、谱展度、谱平坦度 × 2帧", "低"],
        ["EMD包络特征", "4", "Hilbert包络/瞬时频率均值/标准差", "低"],
    ],
    header_color="4472C4"
)
add_caption("表3  完整特征集（共34维）及跨天稳定性评估")

# ================================================================
# 四、前期实验：CFO+KNN方案
# ================================================================
add_heading("四、前期实验：基于连续双帧CFO的KNN分类方案", level=1)

add_para("在项目前期，项目组首先提出了基于连续双帧信号下CFO特征与KNN分类器的设备识别方案。"
         "该方案充分利用了连续帧中频偏信息所携带的设备硬件缺陷特征，"
         "通过简单高效的KNN算法实现设备身份的快速判别。")

add_heading("4.1 方案设计思路", level=2)
add_para("传统方法多基于单帧瞬时CFO，然而单帧特征容易受环境变化和信道时变性的影响。"
         "本方案创新性地提出基于连续双帧信号的CFO特征提取方法，"
         "利用两帧信号间CFO的差值和均值等派生特征，有效抑制外部信道干扰带来的随机波动，"
         "增强设备硬件指纹的稳定呈现。")
add_para("KNN算法的非参数性质使其在样本数量适中、特征维度相对可控的任务中具有很强的实用性。"
         "由于CFO是低维且数值分布清晰的特征，KNN在进行邻近投票时能够准确地将测试样本归类到正确的设备类别中。")

add_heading("4.2 KNN方案实验结果", level=2)
add_para("在最初的8台设备（Heltec LoRa32）上，CFO+KNN方案取得了优异的实验结果：")

make_table(
    ["测试场景", "设备数", "方案", "识别准确率"],
    [
        ["同天测试", "8", "连续双帧CFO + KNN", "99.8%"],
        ["跨天测试", "8", "连续双帧CFO + KNN", "93.0%"],
        ["同天测试", "26 (两种型号)", "连续双帧CFO + KNN", "93.01%"],
        ["跨天测试", "26 (两种型号)", "连续双帧CFO + KNN", "87.61%"],
    ]
)
add_caption("表4  CFO+KNN方案识别结果")

add_para("")
add_para("实验表明连续双帧CFO特征具有良好的跨时间稳定性。"
         "进一步扩大到26台设备（两种型号混合）后，即使训练集和测试集的采集环境存在较大差异，"
         "准确率仍处于较好水平。该实验验证了CFO作为设备指纹核心特征的可行性，"
         "为后续引入多分类器集成学习奠定了基础。")

# ================================================================
# 五、中期实验：CNN频谱图与深度学习探索
# ================================================================
add_heading("五、中期实验：CNN频谱图与深度学习方法探索", level=1)

add_para("在确定特征工程+KNN的基线方案后，项目组在中期阶段对深度学习方法进行了系统性探索，"
         "包括基于CNN的频谱图分类、以及两篇代表性论文（DR-RFF和DeepCRF）的复现与分析。")

add_heading("5.1 CNN频谱图分类实验", level=2)
add_para("项目组尝试将LoRa信号的时域IQ样本转换为STFT频谱图，"
         "然后利用卷积神经网络（CNN）进行端到端分类。"
         "参考Shen等人的工作，CFO补偿后的信号通过短时傅里叶变换生成二维时频图像，"
         "作为CNN的输入。")

add_para("CNN模型架构：采用5层卷积网络（Conv2d + BatchNorm2d + ReLU + MaxPool2d），"
         "通道数逐层递增（8→16→32），经过Dropout正则化后接全连接层（512→N_classes）。"
         "使用Adam优化器、交叉熵损失函数和学习率衰减策略进行训练。"
         "该模型针对18台LoRa设备进行了频谱图分类实验。")

add_para("此外，项目组还实现了基于Triplet Loss的特征学习方案，"
         "通过三元组网络学习设备的嵌入表示，结合KMeans聚类和LocalOutlierFactor（LOF）"
         "进行异常设备检测，探索了开集识别的可能性。")

add_heading("5.2 DR-RFF论文复现", level=2)
add_para("DR-RFF（Disentangled Representation Learning for RF Fingerprint Extraction Under Unknown Channel Statistics）"
         "针对深度学习射频指纹识别中易出现的信道过拟合问题，"
         "提出了基于解耦表示学习的射频特征提取框架。")

add_para("核心方法：通过对抗式训练将接收信号因子化为两部分——"
         "设备相关表示z（真实RFF，随信道不变）与设备无关表示q（背景，包含衰落、多径与噪声）。"
         "随后跨样本随机重组z与q，生成合成信号作为隐式数据增强，"
         "迫使模型仅保留与设备相关的稳健特征。")

add_para("框架由三个模块组成：(1) RFF提取器F——从信号中提取设备相关特征，"
         "通过超球面投影（HP）损失使同一设备特征紧密聚类；"
         "(2) 背景提取器Q——提取设备无关的信道背景特征，通过对抗训练抑制设备信息泄露；"
         "(3) 信号生成器G——将RFF特征与背景信号重组生成合成信号用于数据增强。"
         "训练采用F-step与Q/G-step交替优化策略。")

add_para("复现结果：项目组使用PyTorch实现了DR-RFF的完整训练流程，"
         "在ZigBee和LoRa数据集上验证了其跨信道泛化能力。"
         "实验证实解耦表示学习能有效分离信道效应与设备指纹，"
         "但模型计算开销较大，且需要大量训练数据。")

add_heading("5.3 DeepCRF论文复现", level=2)
add_para("DeepCRF（Deep Learning-Enhanced CSI-Based RF Fingerprinting for Channel-Resilient WiFi Device Identification）"
         "提出了基于深度学习的CSI射频指纹框架，利用WiFi设备CSI数据中的微小硬件特征"
         "实现高鲁棒性的设备识别。")

add_para("核心策略：(1) 模型启发的数据增强（DA）——将去噪后的真实CSI与合成信道/噪声组合，"
         "生成大量具有信道多样性的训练样本；"
         "(2) 监督对比学习（SCL）——预训练使同一设备CSI特征在表示空间中更紧密，不同设备间距离更大；"
         "(3) 决策融合——利用多个CSI测量的预测概率进行平均/乘积/Borda计数等融合策略。")

add_para("网络结构：使用复数值卷积神经网络（ComplexConv2d + ComplexBatchNorm2d），"
         "包含Self-Attention模块和多尺度卷积（1×1、1×3、1×5），仅约12.5万参数。"
         "分两阶段训练：第一阶段对比学习损失训练特征提取器，第二阶段交叉熵损失微调分类器。")

add_para("复现结果：项目组在多种信道条件（室内、室外、移动NLoS）下对9-30台LoRa设备进行了测试，"
         "验证了数据增强和对比学习对信道鲁棒性的提升效果。"
         "实验中采用了多种决策融合策略（多数投票、平均概率、乘积概率、Borda计数），"
         "其中平均概率融合在大多数场景中取得最优结果。")

add_heading("5.4 深度学习方法对比与启示", level=2)

make_table(
    ["方法", "输入形式", "网络结构", "优势", "局限性"],
    [
        ["CNN频谱图", "STFT频谱图", "5层CNN", "端到端学习", "跨天泛化差"],
        ["DR-RFF", "原始IQ", "U-Net+ArcFace", "信道解耦", "计算开销大"],
        ["DeepCRF", "CSI复数值", "复数CNN+注意力", "信道鲁棒", "需大量数据增强"],
        ["Triplet+KNN", "原始IQ", "三元组网络", "开集识别", "训练不稳定"],
    ]
)
add_caption("表5  深度学习方法对比")

add_para("")
add_para("综合分析：深度学习方法虽然在同天场景下能取得较高准确率，"
         "但跨天泛化能力普遍不足，且对计算资源和训练数据量要求较高。"
         "DR-RFF的解耦思想和DeepCRF的数据增强策略为后续工作提供了重要启发——"
         "项目后期可考虑将解耦机制引入特征工程流程，或利用模型驱动的信道仿真扩充训练集。"
         "基于实验对比，项目组最终选择了'特征工程+集成学习'的技术路线，"
         "在保证跨天鲁棒性的同时兼顾计算效率和可解释性。")

# ================================================================
# 六、集成学习框架设计
# ================================================================
add_heading("六、集成学习框架设计与实现", level=1)

add_para("基于前期KNN方案的经验和深度学习方法的对比分析，"
         "项目组设计了动态加权集成学习框架（DualFrameCFOEnsemble），"
         "融合四种互补的基分类器，通过验证集驱动的权重优化实现鲁棒识别。")

add_heading("6.1 基分类器配置", level=2)

make_table(
    ["基分类器", "核心参数", "设计考量"],
    [
        ["SVM (RBF核)", "C=10, Platt概率校准", "非线性决策边界，高维特征空间有效"],
        ["随机森林", "200棵树, n_jobs=-1", "抗过拟合，可评估特征重要性"],
        ["LDA (SVD求解)", "SVD solver", "线性判别，计算开销低，作为线性基线"],
        ["KNN", "k=5, 欧氏距离", "局部相似性捕获，对CFO特征天然适配"],
    ],
    header_color="4472C4"
)
add_caption("表6  集成学习框架基分类器配置")

add_heading("6.2 动态加权融合机制", level=2)
add_para("权重学习：使用验证集（20%训练数据）评估各分类器准确率A_k，"
         "静态权重为 α_k^0 = A_k / ΣA_j。")
add_para("动态置信度：预测时计算各分类器输出概率的信息熵H_k，"
         "动态置信度 β_k = exp(-γ·H_k)，最终权重 α_k = (α_k^0 · β_k) / Σ(α_j^0 · β_j)。"
         "温度参数γ=1.0控制熵的影响强度。")
add_para("双阈值拒识：准确率低于τ₁=0.3的分类器降权，低于τ₂=0.1的分类器权重置零，"
         "确保弱分类器不干扰最终决策，同时为开集识别提供拒识能力。")

# ================================================================
# 七、8台设备消融实验
# ================================================================
add_heading("七、特征消融实验——8台设备（数据集A）", level=1)

add_para("对8台设备（Heltec LoRa32, ID: 009-012, 020-023）开展11种特征组合 × 5种分类器的系统消融实验，"
         "评估各特征对识别性能的贡献。")

with open(os.path.join(RESULTS_DIR, "results_8dev_ablation.json"), "r", encoding="utf-8") as f:
    data_8dev = json.load(f)

headers_abl = ["特征组合", "XGBoost", "随机森林", "SVM", "KNN", "集成学习", "最大正确率"]
col_widths_8 = [Cm(4.5), Cm(1.8), Cm(1.8), Cm(1.6), Cm(1.6), Cm(1.8), Cm(2.0)]

add_heading("7.1 同天测试结果", level=2)
build_ablation_table(data_8dev["table2_same_day"], headers_abl, headers_abl, col_widths_8)
add_caption("表7  8台设备——训练集和测试集在同一天的识别结果")

add_para("")
add_para("同天场景下，CFO特征单独即可达到99.9%（SVM），而相位特征单独仅36%~44%。"
         "Base方案（相位+CFO）达99.8%，RSSI和IQ特征区分度有限（53%~59%）。"
         "集成学习在Base+STFT组合中达到99.9%，与最优单分类器持平。")

add_heading("7.2 跨天测试结果", level=2)
build_ablation_table(data_8dev["table3_another_day"], headers_abl, headers_abl, col_widths_8)
add_caption("表8  8台设备——训练集和测试集在不同天的识别结果")

add_para("")
add_para("跨天场景下RSSI（17.7%）、IQ（20.9%）、单/双帧相位（16.9%/18.4%）几乎完全失效，"
         "证实这些特征不具备跨时间稳定性。CFO特征保持95.8%。"
         "集成学习在Base方案中将最佳单分类器93.0%提升至95.96%，在双帧相位+CFO组合中达96.63%。")

# 图片
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "table2_8dev_same_day.png"), width=Cm(16))
add_caption("图1  8台设备同天测试结果可视化")
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "table3_8dev_another_day.png"), width=Cm(16))
add_caption("图2  8台设备跨天测试结果可视化")

# ================================================================
# 八、18台设备消融实验
# ================================================================
add_heading("八、特征消融实验——18台设备（数据集B）", level=1)

add_para("对数据集B的18台RA01-SC设备开展7种特征组合的消融实验。"
         "该数据集的CFO特征使用frame_info中的m_cfo元数据（范围-1~-6 Hz），"
         "解决了原始提取中CFO补偿残差接近零的问题。")

with open(os.path.join(RESULTS_DIR, "results_18_ablation.json"), "r", encoding="utf-8") as f:
    data_18 = json.load(f)

headers_18 = ["特征组合", "XGBoost", "随机森林", "SVM", "KNN", "集成学习", "最大正确率"]
col_widths_18 = [Cm(4.5), Cm(1.8), Cm(1.8), Cm(1.6), Cm(1.6), Cm(1.8), Cm(2.0)]

add_heading("8.1 同天测试结果", level=2)
build_ablation_table(data_18["table2_same_day"], headers_18, headers_18, col_widths_18)
add_caption("表9  18台设备——训练集和测试集在同一天的识别结果")

add_para("")
add_para("18台设备识别难度较8台设备显著上升。"
         "CFO特征仍是核心（集成学习94.18%），但相位特征区分度更低（23%~38%）。"
         "随机森林在Base方案和双帧相位+CFO组合中均达94.29%。")

add_heading("8.2 跨天测试结果", level=2)
build_ablation_table(data_18["table3_another_day"], headers_18, headers_18, col_widths_18)
add_caption("表10  18台设备——训练集和测试集在不同天的识别结果")

add_para("")
add_para("跨天场景下18台设备识别难度进一步加大。CFO特征保持84.52%（SVM），"
         "双帧相位+CFO组合通过随机森林达85.70%。"
         "单帧/双帧相位在跨天条件下几乎无效（24.7%/37.5%）。")

# 图片
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "table2_same_day.png"), width=Cm(16))
add_caption("图3  18台设备同天测试结果可视化")
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "table3_another_day.png"), width=Cm(16))
add_caption("图4  18台设备跨天测试结果可视化")

# ================================================================
# 九、56台设备集成识别
# ================================================================
add_heading("九、56台设备分层集成识别结果", level=1)

add_para("将三个数据集（8+18+30=56台设备）通过分层集成框架进行统一识别。"
         "每个子数据集独立训练集成学习模型，最终通过块对角混淆矩阵评估整体性能。"
         "该方案无需将所有设备放入同一分类器，有效缓解了大规模分类问题中的类别不平衡和特征分布差异。")

r56_same = {"SVM": 0.8675, "RF": 0.9854, "LDA": 0.9479, "KNN": 0.8184, "集成": 0.9823}
r56_cross = {"SVM": 0.8084, "RF": 0.9632, "LDA": 0.9357, "KNN": 0.7504, "集成": 0.9613}

t56 = doc.add_table(rows=3, cols=6)
t56.alignment = WD_TABLE_ALIGNMENT.CENTER
t56.style = "Table Grid"
t56_headers = ["测试场景", "SVM", "随机森林", "LDA", "KNN", "集成学习"]
for j, h in enumerate(t56_headers):
    cell = t56.rows[0].cells[j]
    set_cell_shading(cell, "2E5A3E")
    set_cell(cell, h, bold=True, font_size=Pt(10),
             color=RGBColor(0xFF, 0xFF, 0xFF))

set_cell(t56.rows[1].cells[0], "同天测试", font_size=Pt(10), bold=True)
for j, (k, v) in enumerate(r56_same.items()):
    cell = t56.rows[1].cells[j + 1]
    is_max = (k == "RF")
    if k == "集成":
        set_cell_shading(cell, "E8F5E9")
    set_cell(cell, f"{v*100:.2f}%", bold=is_max, font_size=Pt(10))

set_cell(t56.rows[2].cells[0], "跨天测试", font_size=Pt(10), bold=True)
for j, (k, v) in enumerate(r56_cross.items()):
    cell = t56.rows[2].cells[j + 1]
    is_max = (k == "RF")
    if k == "集成":
        set_cell_shading(cell, "E8F5E9")
    set_cell(cell, f"{v*100:.2f}%", bold=is_max, font_size=Pt(10))

add_caption("表11  56台设备整体识别正确率")

add_para("")
add_para("同天测试最高98.54%（随机森林），集成学习98.23%；"
         "跨天测试最高96.32%（随机森林），集成学习96.13%。"
         "集成学习在跨天场景中表现出良好的鲁棒性，接近最优单分类器水平，"
         "验证了分层集成架构在多数据集、多型号设备识别中的实用性。")

# 混淆矩阵和对比图
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "cm_56_same_day.png"), width=Cm(14))
add_caption("图5  56台设备同天测试混淆矩阵")
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "cm_56_another_day.png"), width=Cm(14))
add_caption("图6  56台设备跨天测试混淆矩阵")
add_para("")
add_image_safe(os.path.join(RESULTS_DIR, "accuracy_56_comparison.png"), width=Cm(14))
add_caption("图7  56台设备各分类器正确率对比")

# ================================================================
# 十、关键问题排查
# ================================================================
add_heading("十、关键问题排查与解决记录", level=1)

add_heading("10.1 18台设备识别率异常低", level=2)
add_para("问题现象：18台设备同天识别率仅64%，跨天仅17%，远低于预期和8台设备的结果。")
add_para("排查过程：(1) 检查CSV特征文件，发现CFO列数值接近零（~1e-7 Hz）；"
         "(2) 分析代码逻辑，确认cfo_compensation()返回的是补偿后残差而非设备原始CFO；"
         "(3) 查看原始数据frame_info文件，发现m_cfo字段记录了真实CFO（范围-1~-6 Hz）。")
add_para("解决方案：重新编写特征提取脚本(extract_18dev.py)，"
         "使用frame_info中的m_cfo作为CFO特征，仅用计算CFO进行相位补偿。")
add_para("效果：同天64%→90.88%，跨天17%→82.26%，56台整体跨天80.80%→96.13%。"
         "该问题的解决是56台设备跨天性能大幅提升的关键因素。")

add_heading("10.2 同步算法性能优化", level=2)
add_para("问题现象：原始Python循环实现的Schmidl-Cox同步算法处理速度约37秒/设备，"
         "提取18台设备数据需要约11分钟。")
add_para("解决方案：利用numpy的cumsum实现滑动窗口求和，"
         "stride_tricks实现互相关的向量化计算，处理速度提升至约2秒/设备，加速约18倍。")

add_heading("10.3 跨天识别率下降机理分析", level=2)
add_para("实验发现RSSI、IQ统计、单帧/双帧相位等特征在跨天条件下几乎完全失效"
         "（8台设备：13%~21%），原因是这些特征受传播环境（多径、遮挡）、温度变化、"
         "设备老化等因素影响显著，不具备跨时间稳定性。"
         "而CFO特征反映设备本振固有偏差，具有跨时间稳定性（8台设备跨天95.8%）。"
         "STFT和EMD频域特征在跨天场景下同样严重退化（Base+STFT跨天仅58.2%），"
         "因为时频变换容易放大环境噪声的影响。")

# ================================================================
# 十一、实验结论
# ================================================================
add_heading("十一、实验结论与总结", level=1)

add_para("(1) CFO是LoRa设备射频指纹的核心特征，在同天和跨天场景下均保持高区分度。"
         "连续双帧CFO特征相比单帧CFO具有更好的抗干扰能力和跨时间一致性。"
         "相位差分特征作为辅助与CFO组合可进一步提升识别率。")

add_para("(2) RSSI、IQ统计等特征受环境影响大，不适合作为跨时间的设备指纹。"
         "STFT和EMD频域特征在跨天场景下同样存在显著退化，"
         "极大增加了特征维度的同时并未带来性能提升。")

add_para("(3) 深度学习方法（CNN频谱图、DR-RFF、DeepCRF）在同天场景下能取得较高准确率，"
         "但跨天泛化能力不足，且对计算资源和训练数据量要求较高。"
         "DR-RFF的解耦思想和DeepCRF的数据增强策略为未来的混合模型设计提供了重要参考。")

add_para("(4) 集成学习框架通过动态加权融合多分类器，在绝大多数场景中取得最优或接近最优结果，"
         "特别是在跨天测试中优势显著（8台设备Base方案：单分类器最高93.0%→集成95.96%）。"
         "KNN在CFO特征上表现出天然的适配性，作为集成框架中的重要组成部分。")

add_para("(5) 56台设备分层集成方案可行，同天98.23%、跨天96.13%的整体识别率"
         "验证了射频指纹技术在中等规模物联网设备认证中的实用性。"
         "该方案无需将所有设备放入同一分类器，具有良好的可扩展性。")


# ============ 保存 ============
doc.save(OUTPUT_PATH)
print(f"已生成: {OUTPUT_PATH}")
