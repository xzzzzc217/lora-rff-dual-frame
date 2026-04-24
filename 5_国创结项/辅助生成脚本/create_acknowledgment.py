"""生成结题报告致谢 Word 文档"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "致谢.docx")

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(3.0)
section.right_margin = Cm(3.0)
section.top_margin = Cm(3.0)
section.bottom_margin = Cm(3.0)


def add_para(text, bold=False, font_size=Pt(12), alignment=WD_ALIGN_PARAGRAPH.LEFT,
             first_line=True, space_before=Pt(0), space_after=Pt(8)):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = Pt(26)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
    run.font.size = font_size
    run.font.bold = bold
    return p


# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(24)
run = title.add_run("致  谢")
run.font.name = "黑体"
run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
run.font.size = Pt(16)
run.font.bold = True

# 正文
add_para(
    "本项目的顺利完成，离不开众多人士的关心、支持与帮助，在此谨致以诚挚的谢意。"
)

add_para(
    "首先，衷心感谢指导教师吴文甲老师。吴老师在项目立项之初便为我们明确了研究方向，"
    "在整个研究过程中给予了持续的悉心指导。无论是在文献调研、信号处理算法设计，"
    "还是在实验方案的反复讨论与结果分析中，吴老师都不厌其烦地为我们答疑解惑，"
    "指出思路上的偏差，并鼓励我们勇于探索。吴老师严谨的学术态度和开阔的科研视野，"
    "使我们受益匪浅，是我们最重要的引路人。"
)

add_para(
    "感谢东南大学计算机科学与工程学院（计软智学院）提供的实验室场地与软硬件资源支持。"
    "实验室的USRP N210软件无线电平台为本项目的射频信号采集工作提供了关键的硬件保障，"
    "使得多设备、多场景、多天的大规模数据采集成为可能。"
)

add_para(
    "感谢国家大学生创新创业训练计划（国家级项目，编号202510286066）的立项资助。"
    "正是这一平台为项目组提供了宝贵的资金支持和展示机会，"
    "使我们得以在本科阶段系统开展具有一定深度的科研训练。"
)

add_para(
    "特别感谢实验室博士在读学长陈梓超同学对本项目的无私帮助。"
    "在数据采集阶段，陈学长全程带领项目组完成了从USRP设备配置、GNU Radio流程搭建，"
    "到多设备多场景信号采集的全部流程，手把手指导我们应对各种采集中的实际问题，"
    "使我们在较短时间内建立起了完备的射频指纹数据库。"
    "此外，陈学长还在报告撰写和论文修改过程中给予了大量细致的意见与建议，"
    "从内容逻辑到表达规范，均帮助我们提升了文章的质量。"
    "陈学长的悉心指导和慷慨付出，是本项目得以顺利推进的重要支柱，在此深表感谢。"
)

add_para(
    "感谢项目组全体成员——刘博文、赵子辰、昝召朋同学的紧密协作。"
    "项目历时一年，期间我们共同完成了文献调研、系统设计、数据采集、代码开发、"
    "实验对比与报告撰写等各项工作。每一次讨论、每一次调试、每一次结果分析，"
    "都凝聚着大家的智慧与努力。团队的协作精神和互相鼓励是我们克服困难、"
    "持续推进项目的重要动力。"
)

add_para(
    "本项目在研究过程中参考了国内外众多学者的研究成果，"
    "包括Shen等人在LoRa频谱图CNN识别方面的工作、"
    "DR-RFF和DeepCRF等深度学习框架的原创论文，"
    "以及Schmidl-Cox同步算法和Robyns等人精同步方案的理论基础。"
    "在此对上述研究者的开创性贡献表示感谢。"
)

add_para(
    "此外，感谢所有在实验数据采集过程中给予帮助的同学，"
    "以及在报告写作阶段提出宝贵意见的朋友们。"
    "你们的支持是本项目能够顺利完成的重要保障。"
)

add_para(
    "本项目从CFO+KNN的初步方案，到深度学习框架的探索与复现，"
    "再到最终集成学习框架在56台设备上的验证，"
    "每一步都凝结着项目组的反复思考与实践。"
    "研究过程中我们遇到了诸多困难——"
    "同步算法的性能瓶颈、CFO特征提取的数据问题、跨天场景下的泛化挑战，"
    "每一次解决问题的过程都是一次深刻的学习经历，"
    "让我们对射频指纹识别这一领域有了更扎实、更全面的认识。"
    "这段科研经历将成为我们宝贵的财富。"
)

add_para(
    "最后，感谢每一位关注和支持本项目的老师与同学。"
    "由于项目组成员能力和经验尚有不足，研究成果难免存在局限与不完善之处，"
    "敬请批评指正。我们将在未来的学习与研究中继续努力，不断深化和完善相关工作。"
)

# 签名
sig = doc.add_paragraph()
sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
sig.paragraph_format.space_before = Pt(24)
sig.paragraph_format.space_after = Pt(4)
sig.paragraph_format.first_line_indent = Pt(0)
run = sig.add_run("项目组全体成员")
run.font.name = "宋体"
run._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
run.font.size = Pt(12)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
date_p.paragraph_format.space_before = Pt(0)
date_p.paragraph_format.first_line_indent = Pt(0)
run = date_p.add_run("2025年10月")
run.font.name = "宋体"
run._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
run.font.size = Pt(12)

doc.save(OUTPUT_PATH)
print(f"已生成: {OUTPUT_PATH}")
