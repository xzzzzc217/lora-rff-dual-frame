"""生成8台设备消融实验结果的Word文档（表2+表3）"""

import json
import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

RESULTS_DIR = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = os.path.join(RESULTS_DIR, "results_8dev_ablation.json")
OUTPUT_PATH = os.path.join(RESULTS_DIR, "8台设备消融实验结果.docx")

with open(JSON_PATH, "r", encoding="utf-8") as f:
    data = json.load(f)

headers = ["特征组合", "XGBoost", "随机森林", "SVM", "KNN", "集成学习", "最大正确率"]
col_keys = ["特征组合", "XGBoost", "随机森林", "SVM", "KNN", "集成学习", "最大正确率"]


def fmt_pct(val_str):
    v = float(val_str)
    return f"{v*100:.1f}%"


def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color,
    })
    shading.append(shading_elem)


def set_cell_text(cell, text, bold=False, font_name="宋体", font_size=Pt(10.5),
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    """设置单元格文本"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def build_table(doc, table_data, title):
    """构建一个带标题的消融实验表格"""
    # 标题
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(12)
    title_para.paragraph_format.space_after = Pt(8)
    run = title_para.add_run(title)
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    run.font.size = Pt(12)
    run.font.bold = True

    n_rows = len(table_data) + 1  # +1 for header
    n_cols = len(headers)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 设置列宽
    col_widths = [Cm(5.0), Cm(2.0), Cm(2.0), Cm(1.8), Cm(1.8), Cm(2.0), Cm(2.2)]
    for i, width in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = width

    # 表头
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, "2E5A3E")
        set_cell_text(cell, header, bold=True, font_name="宋体",
                      font_size=Pt(10), color=RGBColor(0xFF, 0xFF, 0xFF))

    # 数据行
    for i, row_data in enumerate(table_data):
        max_val = float(row_data["最大正确率"])

        for j, key in enumerate(col_keys):
            cell = table.rows[i + 1].cells[j]
            val = row_data[key]

            if key == "特征组合":
                set_cell_text(cell, val, bold=False, font_size=Pt(10),
                              alignment=WD_ALIGN_PARAGRAPH.LEFT)
            else:
                text = fmt_pct(val)
                is_max = abs(float(val) - max_val) < 0.0001
                is_ensemble = (key == "集成学习")

                if is_ensemble:
                    set_cell_shading(cell, "E8F5E9")

                set_cell_text(cell, text, bold=is_max, font_size=Pt(10))

    return table


# 创建文档
doc = Document()

# 设置页面为A4
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

# 表2: 同天
build_table(doc, data["table2_same_day"],
            "表2  训练集和测试集在同一天的识别结果")

# 空行
doc.add_paragraph()

# 表3: 跨天
build_table(doc, data["table3_another_day"],
            "表3  训练集和测试集在不同天的识别结果")

doc.save(OUTPUT_PATH)
print(f"已生成: {OUTPUT_PATH}")
